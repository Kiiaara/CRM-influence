"""База блогеров (импорт из многолистовой таблицы), тип участника сделки и ссылка на КП."""
import asyncio
import io
from datetime import datetime, timedelta

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

import blogger_sheet
import main
from database import Base, engine, SessionLocal
from models.advertiser import Advertiser
from models.auth_session import AuthSession
from models.user import User
from models.workspace import Workspace, WorkspaceMember

ADMIN_TG = 105001
VIEWER_TG = 105002


def _workbook_bytes() -> bytes:
    """Как в реальной таблице: листы по площадкам, заголовки называются по-разному,
    на одном листе над заголовком есть служебная строка, имя - гиперссылка на канал."""
    wb = Workbook()
    yt = wb.active
    yt.title = "YouTube"
    yt.append(["Прайс на сентябрь"])
    yt.append(["Блогер", "ТГ для связи", "Подписчики", "Стоимость интеграции", "Тематика"])
    yt.append(["Иван Игры", "@ivan_games", "1,2 млн", "150 000", "игры"])
    yt.cell(row=3, column=1).hyperlink = "https://youtube.com/@ivangames"
    yt.append([None, None, None, None, None])
    yt.append(["Маша Тех", "", "350000", "80000", "техника"])

    ig = wb.create_sheet("Instagram")
    ig.append(["Ссылка на аккаунт", "Telegram", "Охват reels"])
    ig.append(["https://instagram.com/katya.style", "https://t.me/katya", "200к"])

    wb.create_sheet("Пустой лист")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_parse_workbook_reads_all_sheets():
    items = blogger_sheet.parse_workbook(_workbook_bytes())
    by_name = {i["name"]: i for i in items}
    assert set(by_name) == {"Иван Игры", "Маша Тех", "katya.style"}

    ivan = by_name["Иван Игры"]
    assert ivan["platform"] == "YouTube"
    assert ivan["telegram"] == "@ivan_games"
    assert ivan["url"] == "https://youtube.com/@ivangames"
    assert ivan["subscribers"] == 1_200_000
    assert ivan["price"] == 150_000
    assert ivan["category"] == "игры"

    katya = by_name["katya.style"]
    assert katya["platform"] == "Instagram"
    assert katya["url"] == "https://instagram.com/katya.style"
    assert katya["telegram"] == "https://t.me/katya"
    assert katya["avg_views"] == 200_000


def test_export_url_from_sheet_link():
    url = "https://docs.google.com/spreadsheets/d/18js3Qu_VJ0-S0J/edit?gid=216813285#gid=216813285"
    assert blogger_sheet.export_url(url) == "https://docs.google.com/spreadsheets/d/18js3Qu_VJ0-S0J/export?format=xlsx"
    with pytest.raises(blogger_sheet.SheetImportError):
        blogger_sheet.export_url("https://example.com/table")


@pytest.fixture(scope="module")
def ctx():
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    db.add_all([
        User(tg_id=ADMIN_TG, role="admin", label="Админ"),
        User(tg_id=VIEWER_TG, role="viewer", label="Зритель"),
    ])
    db.commit()
    ws = Workspace(title="Блогеры", owner_tg_id=ADMIN_TG)
    db.add(ws)
    db.commit()
    db.refresh(ws)
    now = datetime.now()
    db.add_all([
        WorkspaceMember(workspace_id=ws.id, user_tg_id=ADMIN_TG, role="owner", joined_at=now),
        WorkspaceMember(workspace_id=ws.id, user_tg_id=VIEWER_TG, role="viewer", joined_at=now),
        AuthSession(token="blog-admin", tg_id=ADMIN_TG, expires_at=now + timedelta(days=1)),
        AuthSession(token="blog-viewer", tg_id=VIEWER_TG, expires_at=now + timedelta(days=1)),
    ])
    adv = Advertiser(workspace_id=ws.id, name="Бренд КП")
    db.add(adv)
    db.commit()
    data = {"ws_id": ws.id, "adv_id": adv.id}
    db.close()
    with TestClient(main.app) as client:
        yield client, data


def _as(client, token: str, ws_id: int):
    client.cookies.set("zametochnitsa_session", token)
    return {"X-Workspace-Id": str(ws_id)}


def test_import_upserts_and_keeps_manual_values(ctx):
    client, data = ctx
    h = _as(client, "blog-admin", data["ws_id"])
    files = {"file": ("bloggers.xlsx", _workbook_bytes(), "application/octet-stream")}
    r = client.post("/api/blogger-profiles/import", files=files, headers=h)
    assert r.status_code == 200, r.text
    assert r.json() == {"created": 3, "updated": 0, "deleted": 0, "total": 3}

    masha = next(b for b in client.get("/api/blogger-profiles", headers=h).json() if b["name"] == "Маша Тех")
    r = client.patch(f"/api/blogger-profiles/{masha['id']}", json={"telegram": "@masha"}, headers=h)
    assert r.status_code == 200

    # повторный импорт: те же люди обновляются, а пустая ячейка TG не затирает ручную правку
    r = client.post("/api/blogger-profiles/import", files=files, headers=h)
    assert r.json() == {"created": 0, "updated": 3, "deleted": 0, "total": 3}
    masha = next(b for b in client.get("/api/blogger-profiles", headers=h).json() if b["name"] == "Маша Тех")
    assert masha["telegram"] == "@masha"


def test_sheet_url_setting_syncs_tabs_and_columns(ctx, monkeypatch):
    client, data = ctx
    h = _as(client, "blog-admin", data["ws_id"])

    async def fake_download(url):
        return _workbook_bytes()
    monkeypatch.setattr(blogger_sheet, "download_sheet", fake_download)

    bad = client.put("/api/blogger-profiles/sheet", json={"url": "https://example.com"}, headers=h)
    assert bad.status_code == 400
    link = "https://docs.google.com/spreadsheets/d/abc123/edit"
    r = client.put("/api/blogger-profiles/sheet", json={"url": link}, headers=h)
    assert r.status_code == 200
    info = client.get("/api/blogger-profiles/sheet", headers=h).json()
    assert info["url"] == link
    # вкладки - листы с блогерами в порядке таблицы, пустой лист не попадает
    assert info["tabs"] == ["YouTube", "Instagram"]
    assert info["last_sync"]["error"] is None

    ivan = next(b for b in client.get("/api/blogger-profiles", headers=h).json() if b["name"] == "Иван Игры")
    assert ivan["source"] == "sheet"
    # все колонки листа как в таблице, по порядку, со ссылкой из гиперссылки
    assert [c["h"] for c in ivan["sheet_row"]] == ["Блогер", "ТГ для связи", "Подписчики", "Стоимость интеграции", "Тематика"]
    assert ivan["sheet_row"][0] == {"h": "Блогер", "v": "Иван Игры", "u": "https://youtube.com/@ivangames"}

    hv = _as(client, "blog-viewer", data["ws_id"])
    assert client.put("/api/blogger-profiles/sheet", json={"url": ""}, headers=hv).status_code == 403


def test_deal_kp_link_and_blogger_participant(ctx):
    client, data = ctx
    h = _as(client, "blog-admin", data["ws_id"])
    kp = "https://docs.google.com/spreadsheets/d/kp1/edit"
    r = client.post("/api/integrations", json={"advertiser_id": data["adv_id"], "kp_sheet_url": kp}, headers=h)
    assert r.status_code == 201
    it = r.json()
    assert it["kp_sheet_url"] == kp

    r = client.patch(f"/api/integrations/{it['id']}", json={"kp_sheet_url": " https://docs.google.com/spreadsheets/d/kp2 "}, headers=h)
    assert r.json()["kp_sheet_url"] == "https://docs.google.com/spreadsheets/d/kp2"

    r = client.post(f"/api/integrations/{it['id']}/streamers", json={"streamer_name": "Иван Игры", "talent_type": "blogger"}, headers=h)
    assert r.status_code == 201
    assert r.json()["talent_type"] == "blogger"

    # по умолчанию - стример, а мусор в типе не принимаем
    r = client.post(f"/api/integrations/{it['id']}/streamers", json={"streamer_name": "Стример"}, headers=h)
    assert r.json()["talent_type"] == "streamer"
    r = client.patch(f"/api/integrations/streamers/{r.json()['id']}", json={"talent_type": "robot"}, headers=h)
    assert r.status_code == 400


def test_sync_removes_rows_deleted_from_sheet_but_keeps_manual(ctx, monkeypatch):
    client, data = ctx
    h = _as(client, "blog-admin", data["ws_id"])
    manual = client.post("/api/blogger-profiles", json={"name": "Ручной", "platform": "YouTube"}, headers=h)
    assert manual.status_code == 201

    # из листа YouTube убрали Машу, лист Instagram не прочитался (заголовки сломали)
    wb = Workbook()
    yt = wb.active
    yt.title = "YouTube"
    yt.append(["Блогер", "ТГ для связи"])
    yt.append(["Иван Игры", "@ivan_games"])
    broken = wb.create_sheet("Instagram")
    broken.append(["что-то", "непонятное"])
    broken.append(["x", "y"])
    buf = io.BytesIO()
    wb.save(buf)

    async def fake_download(url):
        return buf.getvalue()
    monkeypatch.setattr(blogger_sheet, "download_sheet", fake_download)

    r = client.post("/api/blogger-profiles/sync", headers=h)
    assert r.status_code == 200, r.text
    assert r.json()["deleted"] == 1
    names = {b["name"] for b in client.get("/api/blogger-profiles", headers=h).json()}
    assert "Маша Тех" not in names  # удалена из таблицы
    assert "katya.style" in names  # её лист не прочитался - не трогаем
    assert "Ручной" in names  # заведён руками


def test_scheduler_job_syncs_sheet(ctx, monkeypatch):
    import scheduler
    calls = []

    async def fake_sync(db):
        calls.append(True)
        return {}
    monkeypatch.setattr(blogger_sheet, "sync_from_sheet", fake_sync)
    asyncio.run(scheduler._sync_bloggers_sheet())
    assert calls == [True]
